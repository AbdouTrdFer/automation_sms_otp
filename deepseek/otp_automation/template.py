"""Create an example Excel input template with synthetic (non-routable) numbers."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# Fictional reserved ranges: NANP 555-01xx and UK 07700 900xxx.
SYNTHETIC_ROWS = [
    ("+1 202 555 0143", "synthetic - NANP reserved"),
    ("+1 (202) 555-0144", "synthetic - NANP reserved"),
    ("+1 202 555 0143", "duplicate on purpose (will be skipped)"),
    ("+44 7700 900123", "synthetic - UK reserved"),
    ("+1 202 555 0145", "synthetic - NANP reserved"),
    ("not-a-number", "invalid on purpose (will be skipped)"),
    ("", "empty on purpose (will be skipped)"),
    ("+1 202 555 0146", "synthetic - NANP reserved"),
]


def create_template(output: str | Path = "input/test_numbers.xlsx") -> Path:
    path = Path(output)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "numbers"

    headers = ["phone_number", "note"]
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="center")

    for row_idx, (phone, note) in enumerate(SYNTHETIC_ROWS, start=2):
        ws.cell(row=row_idx, column=1, value=phone).number_format = "@"
        ws.cell(row=row_idx, column=2, value=note)

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 40
    ws.freeze_panes = "A2"

    wb.save(path)
    return path


if __name__ == "__main__":
    created = create_template()
    print(f"Created: {created}")