"""Automated reporting module for OTP automation."""

from __future__ import annotations

from collections import Counter
import csv
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import ExecutionResult, SkippedItem, TestStatus


REPORT_HEADERS = [
    "Row",
    "Phone Number",
    "Status",
    "Category",
    "Details",
    "Duration (s)",
    "Timestamp",
]

STATUS_COLORS = {
    TestStatus.SUCCESS.value: "E2EFDA",  # soft green
    TestStatus.FAILED.value: "FCE4D6",   # soft pink/red
    TestStatus.ERROR.value: "F8CBAD",    # soft coral
    TestStatus.UNKNOWN.value: "EDEDED",  # soft light gray
    TestStatus.SKIPPED.value: "FFF2CC",  # soft yellow
}

STATUS_TEXT_COLORS = {
    TestStatus.SUCCESS.value: "276A3C",
    TestStatus.FAILED.value: "C00000",
    TestStatus.ERROR.value: "C65911",
    TestStatus.UNKNOWN.value: "595959",
    TestStatus.SKIPPED.value: "7F6000",
}


def _record_to_dict(item: ExecutionResult | SkippedItem) -> dict[str, Any]:
    if isinstance(item, ExecutionResult):
        raw_number = item.test_case.raw or item.test_case.normalized
        ts = item.timestamp[:19].replace("T", " ") if item.timestamp else ""
        details = item.error or item.message or ("OTP verification sent" if item.status == TestStatus.SUCCESS else "")
        return {
            "Row": item.test_case.row,
            "Phone Number": str(raw_number),
            "Status": item.status.value,
            "Category": item.category or "",
            "Details": details,
            "Duration (s)": round(item.duration_seconds, 2),
            "Timestamp": ts,
        }
    else:  # SkippedItem
        return {
            "Row": item.row,
            "Phone Number": str(item.raw or ""),
            "Status": TestStatus.SKIPPED.value,
            "Category": "SKIPPED",
            "Details": f"Skipped: {item.reason}",
            "Duration (s)": 0.0,
            "Timestamp": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"),
        }


def write_report(
    results: list[ExecutionResult | SkippedItem],
    output_file: str | Path,
    also_write_csv: bool = True,
) -> Path:
    """Generate Excel workbook report and optional CSV report."""
    out_path = Path(output_file)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    records = [_record_to_dict(r) for r in results]

    # Excel Generation
    wb = Workbook()

    # -----------------------------------------------------------------
    # Sheet 1: Detailed Results
    # -----------------------------------------------------------------
    ws = wb.active
    ws.title = "Results"

    header_fill = PatternFill("solid", fgColor="1F4E79")  # Professional navy
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for col_idx, header in enumerate(REPORT_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

    ws.row_dimensions[1].height = 28

    # Populate rows
    for row_idx, rec in enumerate(records, start=2):
        status_val = str(rec["Status"])
        row_fill_color = STATUS_COLORS.get(status_val)
        status_text_color = STATUS_TEXT_COLORS.get(status_val, "000000")
        row_fill = PatternFill("solid", fgColor=row_fill_color) if row_fill_color else None

        for col_idx, header in enumerate(REPORT_HEADERS, start=1):
            val = rec[header]
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill

            if header == "Row":
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif header == "Phone Number":
                cell.number_format = "@"  # treat as text to keep full digits intact
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif header == "Status":
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(bold=True, color=status_text_color)
            elif header == "Category":
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif header == "Duration (s)":
                cell.number_format = "0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif header == "Timestamp":
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[row_idx].height = 22

    # Auto column width
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    ws.freeze_panes = "A2"
    if records:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(REPORT_HEADERS))}{len(records) + 1}"

    # -----------------------------------------------------------------
    # Sheet 2: Executive Summary Dashboard
    # -----------------------------------------------------------------
    ws_summary = wb.create_sheet(title="Summary")
    ws_summary.views.sheetView[0].showGridLines = True

    # Title
    title_cell = ws_summary.cell(row=1, column=1, value="OTP Automation Execution Summary")
    title_cell.font = Font(size=14, bold=True, color="1F4E79")
    ws_summary.row_dimensions[1].height = 28

    total = len(records)
    success_cnt = sum(1 for r in records if r["Status"] == TestStatus.SUCCESS.value)
    failed_cnt = sum(1 for r in records if r["Status"] == TestStatus.FAILED.value)
    error_cnt = sum(1 for r in records if r["Status"] == TestStatus.ERROR.value)
    unknown_cnt = sum(1 for r in records if r["Status"] == TestStatus.UNKNOWN.value)
    skipped_cnt = sum(1 for r in records if r["Status"] == TestStatus.SKIPPED.value)
    success_rate = f"{(success_cnt / total * 100):.1f}%" if total > 0 else "0.0%"

    metrics = [
        ("Total Phone Numbers", total, "EDEDED", "000000"),
        ("SUCCESS (OTP Sent)", success_cnt, STATUS_COLORS[TestStatus.SUCCESS.value], STATUS_TEXT_COLORS[TestStatus.SUCCESS.value]),
        ("FAILED", failed_cnt, STATUS_COLORS[TestStatus.FAILED.value], STATUS_TEXT_COLORS[TestStatus.FAILED.value]),
        ("ERROR", error_cnt, STATUS_COLORS[TestStatus.ERROR.value], STATUS_TEXT_COLORS[TestStatus.ERROR.value]),
        ("UNKNOWN", unknown_cnt, STATUS_COLORS[TestStatus.UNKNOWN.value], STATUS_TEXT_COLORS[TestStatus.UNKNOWN.value]),
        ("SKIPPED", skipped_cnt, STATUS_COLORS[TestStatus.SKIPPED.value], STATUS_TEXT_COLORS[TestStatus.SKIPPED.value]),
        ("Success Rate", success_rate, "D9E1F2", "1F4E79"),
    ]

    ws_summary.cell(row=3, column=1, value="Metric").font = Font(bold=True, color="FFFFFF")
    ws_summary.cell(row=3, column=1).fill = header_fill
    ws_summary.cell(row=3, column=2, value="Value").font = Font(bold=True, color="FFFFFF")
    ws_summary.cell(row=3, column=2).fill = header_fill

    for r_idx, (m_label, m_val, bg, fg) in enumerate(metrics, start=4):
        c1 = ws_summary.cell(row=r_idx, column=1, value=m_label)
        c2 = ws_summary.cell(row=r_idx, column=2, value=m_val)
        c1.border = thin_border
        c2.border = thin_border
        c1.font = Font(bold=(m_label in ("Total Phone Numbers", "Success Rate")))
        c2.font = Font(bold=True, color=fg)
        c2.alignment = Alignment(horizontal="center")
        if bg:
            c2.fill = PatternFill("solid", fgColor=bg)

    ws_summary.column_dimensions["A"].width = 28
    ws_summary.column_dimensions["B"].width = 18

    # Category breakdown table
    cat_row_start = len(metrics) + 6
    ws_summary.cell(row=cat_row_start, column=1, value="Response Categories").font = Font(size=12, bold=True, color="1F4E79")
    ws_summary.cell(row=cat_row_start + 1, column=1, value="Category").font = Font(bold=True, color="FFFFFF")
    ws_summary.cell(row=cat_row_start + 1, column=1).fill = header_fill
    ws_summary.cell(row=cat_row_start + 1, column=2, value="Count").font = Font(bold=True, color="FFFFFF")
    ws_summary.cell(row=cat_row_start + 1, column=2).fill = header_fill

    cats = Counter(r["Category"] for r in records if r["Category"])
    for idx, (cat_name, cat_count) in enumerate(sorted(cats.items(), key=lambda x: -x[1]), start=cat_row_start + 2):
        c1 = ws_summary.cell(row=idx, column=1, value=cat_name)
        c2 = ws_summary.cell(row=idx, column=2, value=cat_count)
        c1.border = thin_border
        c2.border = thin_border
        c2.alignment = Alignment(horizontal="center")

    wb.save(out_path)

    # Optional CSV generation
    if also_write_csv:
        csv_path = out_path.with_suffix(".csv")
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=REPORT_HEADERS)
            writer.writeheader()
            for r in records:
                writer.writerow(r)

    return out_path


def summarize(results: list[ExecutionResult]) -> dict[str, Any]:
    """Calculate counts and aggregate statistics from test execution."""
    counts = {
        TestStatus.SUCCESS.value: 0,
        TestStatus.FAILED.value: 0,
        TestStatus.ERROR.value: 0,
        TestStatus.UNKNOWN.value: 0,
        TestStatus.SKIPPED.value: 0,
    }
    categories: dict[str, int] = {}
    total_duration = 0.0

    for r in results:
        status_str = r.status.value if isinstance(r.status, TestStatus) else str(r.status)
        counts[status_str] = counts.get(status_str, 0) + 1
        total_duration += r.duration_seconds
        if r.category:
            categories[r.category] = categories.get(r.category, 0) + 1

    return {
        "total": len(results),
        "counts": counts,
        "categories": categories,
        "total_duration_seconds": round(total_duration, 2),
    }


def format_summary(summary: dict[str, Any]) -> str:
    """Format execution summary for console output."""
    total = summary["total"]
    counts = summary["counts"]
    dur = summary.get("total_duration_seconds", 0.0)

    lines = [
        "",
        "=" * 50,
        "            OTP AUTOMATION RUN SUMMARY            ",
        "=" * 50,
        f" Total Executed  : {total}",
        f" SUCCESS         : {counts.get('SUCCESS', 0)}",
        f" FAILED          : {counts.get('FAILED', 0)}",
        f" ERROR           : {counts.get('ERROR', 0)}",
        f" UNKNOWN         : {counts.get('UNKNOWN', 0)}",
        f" SKIPPED         : {counts.get('SKIPPED', 0)}",
        f" Total Duration  : {dur:.2f}s",
    ]

    categories = summary.get("categories", {})
    if categories:
        lines.append("-" * 50)
        lines.append(" Response Categories:")
        for cat, cnt in categories.items():
            lines.append(f"   - {cat}: {cnt}")

    lines.append("=" * 50)
    lines.append("")
    return "\n".join(lines)
